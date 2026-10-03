"""Document content extractors for Phase 5.

This module provides deterministic text extraction for supported file types.
It does not contain any database or filesystem security logic; it expects
a validated canonical Path to a readable file.
"""

from pathlib import Path
from typing import Optional, Tuple
import pypdf
import docx


def extract_txt(path: Path) -> Tuple[Optional[str], Optional[int], Optional[int], Optional[str]]:
    """Extract text from a .txt file.

    Returns:
        (text, char_count, page_count, error_info)
    """
    try:
        # Try UTF-8 first
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            # Fallback to latin-1 which rarely raises DecodeError, 
            # though the characters might be garbage if it's actually binary.
            text = path.read_text(encoding="latin-1")
        except Exception as e:
            return None, None, None, f"Failed to decode text file: {e}"
    except Exception as e:
        return None, None, None, f"Failed to read file: {e}"

    if not text.strip():
        return "", 0, None, None

    return text, len(text), None, None


def extract_pdf(path: Path) -> Tuple[Optional[str], Optional[int], Optional[int], Optional[str]]:
    """Extract text from a .pdf file using pypdf.

    Returns:
        (text, char_count, page_count, error_info)
    """
    try:
        text_parts = []
        page_count = 0
        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            
            if reader.is_encrypted:
                return None, None, None, "PDF is encrypted/password-protected"
                
            page_count = len(reader.pages)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
                    
        text = "\n".join(text_parts)
        if not text.strip():
            return "", 0, page_count, None
            
        return text, len(text), page_count, None
    except Exception as e:
        return None, None, None, f"PDF extraction failed: {type(e).__name__} - {e}"


def extract_docx(path: Path) -> Tuple[Optional[str], Optional[int], Optional[int], Optional[str]]:
    """Extract text from a .docx file using python-docx.

    Returns:
        (text, char_count, page_count, error_info)
    """
    try:
        doc = docx.Document(path)
        
        text_parts = []
        for para in doc.paragraphs:
            if para.text:
                text_parts.append(para.text)
                
        # Basic table text extraction (optional but highly requested in document analysis)
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text:
                        row_text.append(cell.text)
                if row_text:
                    text_parts.append(" | ".join(row_text))
                    
        text = "\n".join(text_parts)
        if not text.strip():
            return "", 0, None, None
            
        return text, len(text), None, None
    except Exception as e:
        return None, None, None, f"DOCX extraction failed: {type(e).__name__} - {e}"
